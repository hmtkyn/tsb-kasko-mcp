"""Reader for the monthly kasko spreadsheets published by TSB.

TSB publishes one workbook per month under ``/content/InsuranceExcelFiles``. The
column layout has changed between publications and carries no version marker, so
the reader locates the header row and maps columns by keyword instead of relying
on fixed positions.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from .exceptions import TsbParseError
from .models import ArchiveRow
from .parsing import fold, matches_all_terms, parse_amount, split_terms

_HEADER_SCAN_ROWS = 25

_COLUMN_KEYWORDS: dict[str, tuple[str, ...]] = {
    "brand_code": ("marka kodu", "markakodu", "brand code"),
    "model_code": ("model kodu", "modelkodu", "tip kodu", "model code"),
    "model_year": ("model yili", "modelyili", "model year", "yil"),
    "brand": ("marka", "brand"),
    "model": ("model", "tip", "arac", "vehicle"),
    "amount": ("kasko bedeli", "kaskobedeli", "bedel", "deger", "tutar", "amount", "value"),
}

_FIELD_PRIORITY = ("brand_code", "model_code", "model_year", "amount", "brand", "model")

_COMBINED_FIELD = "brand_model"

_COMBINED_SEPARATORS = (" - ", "-", "/")


def _cell_text(value: Any) -> str:
    """Render a spreadsheet cell as trimmed text.

    Args:
        value: Raw cell value.

    Returns:
        The cell rendered as a string, empty when the cell is blank.
    """
    return "" if value is None else str(value).strip()


def _match_column(header: str) -> str | None:
    """Map a header label onto a logical field name.

    Fields are tested in priority order so that a header such as ``Marka Kodu``
    binds to ``brand_code`` rather than to the broader ``brand`` rule.

    Args:
        header: Header label as written in the workbook.

    Returns:
        The logical field name, or ``None`` when the header is not recognised.
    """
    folded = fold(header)
    if not folded:
        return None
    if "marka" in folded and ("model" in folded or "tip" in folded):
        return _COMBINED_FIELD
    for field in _FIELD_PRIORITY:
        if any(keyword in folded for keyword in _COLUMN_KEYWORDS[field]):
            return field
    return None


def _split_combined(text: str) -> tuple[str | None, str | None]:
    """Split a combined brand and model label into its two parts.

    Older TSB sheets publish a single ``Araç Marka - Model`` column holding a
    value such as ``AUDI-80 2600``. The leading token is the brand and the rest
    is the model designation.

    Args:
        text: Combined label as written in the workbook.

    Returns:
        The brand and model parts. The brand is ``None`` when the label carries
        no recognisable separator.
    """
    if not text:
        return None, None
    for separator in _COMBINED_SEPARATORS:
        if separator in text:
            head, _, tail = text.partition(separator)
            head, tail = head.strip(), tail.strip()
            if head and tail:
                return head, tail
    return None, text


def _locate_header(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    """Find the header row and the column index of each recognised field.

    Args:
        rows: The first rows of the sheet.

    Returns:
        The index of the header row and the mapping of logical field names to
        column indices.

    Raises:
        TsbParseError: When no row looks like a header.
    """
    best_index = -1
    best_mapping: dict[str, int] = {}
    for index, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        mapping: dict[str, int] = {}
        for column, value in enumerate(row):
            field = _match_column(_cell_text(value))
            if field is not None and field not in mapping:
                mapping[field] = column
        if len(mapping) > len(best_mapping):
            best_index, best_mapping = index, mapping
    if best_index < 0 or "amount" not in best_mapping:
        raise TsbParseError(
            "Could not locate a header row carrying a kasko value column in the workbook"
        )
    return best_index, best_mapping


def _as_int(value: Any) -> int | None:
    """Convert a spreadsheet cell into an integer when possible.

    Args:
        value: Raw cell value.

    Returns:
        The integer value, or ``None`` when the cell does not hold one.
    """
    text = _cell_text(value)
    if not text:
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def parse_workbook(
    content: bytes,
    *,
    year: int,
    month: int,
    query: str | None = None,
    limit: int = 100,
) -> list[ArchiveRow]:
    """Read valuation rows out of a monthly kasko workbook.

    Args:
        content: Raw bytes of the downloaded workbook.
        year: Publication year, recorded on every row.
        month: Publication month, recorded on every row.
        query: Optional free text filter applied to the brand and model columns.
        limit: Maximum number of rows returned.

    Returns:
        The matching valuation rows, at most ``limit`` entries.

    Raises:
        TsbParseError: When the workbook cannot be opened or has no usable header.
    """
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise TsbParseError(f"Could not open the TSB archive workbook: {error}") from error

    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    header_index, columns = _locate_header(rows)
    terms = split_terms(query) if query else []
    results: list[ArchiveRow] = []

    for row in rows[header_index + 1 :]:
        if len(results) >= limit:
            break
        if _COMBINED_FIELD in columns:
            brand, model = _split_combined(_cell_text(row[columns[_COMBINED_FIELD]]))
        else:
            brand = _cell_text(row[columns["brand"]]) if "brand" in columns else None
            model = _cell_text(row[columns["model"]]) if "model" in columns else None
        if not brand and not model:
            continue
        if terms and not matches_all_terms(fold(f"{brand or ''} {model or ''}"), terms):
            continue

        amount = parse_amount(row[columns["amount"]])
        if amount <= 0:
            continue

        brand_code = _as_int(row[columns["brand_code"]]) if "brand_code" in columns else None
        model_code = _as_int(row[columns["model_code"]]) if "model_code" in columns else None
        results.append(
            ArchiveRow(
                year=year,
                month=month,
                model_year=_as_int(row[columns["model_year"]]) if "model_year" in columns else None,
                brand=brand or None,
                model=model or None,
                brand_code=brand_code,
                model_code=model_code,
                vehicle_code=(
                    f"{brand_code}-{model_code}"
                    if brand_code is not None and model_code is not None
                    else None
                ),
                amount=amount,
            )
        )
    return results
